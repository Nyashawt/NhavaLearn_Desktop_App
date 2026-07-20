"""
Library documents — admin-uploaded reference material (PDF/Word), scoped by
grade + subject, searched via SQLite FTS5 keyword search to ground AI
generation in the school's own material rather than pure model knowledge.

No embeddings, no vector index, no extra model to manage — FTS5 ships with
SQLite and matches the offline-first / fixed-hardware constraints already
governing local generation in ai_generator.py.
"""
import os
import time
from typing import Dict, List, Optional

from . import db


def extract_text(file_path: str) -> str:
    """Dispatch on extension. Lazy imports — mirrors ai_generator.py's lazy
    `from llama_cpp import Llama`, keeping heavy libs out of module import time."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        import pdfplumber
        text = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text.append(page_text)
        return "\n".join(text)
    if ext == ".docx":
        import docx
        d = docx.Document(file_path)
        return "\n".join(p.text for p in d.paragraphs if p.text)
    raise ValueError(f"Unsupported document type: {ext}")


def store_document(src_path: str, title: str, subject_id: Optional[int], grade: str) -> Dict:
    ext = os.path.splitext(src_path)[1].lower()
    safe_name = f"doc_{int(time.time())}_{os.path.basename(src_path)}"
    dest = os.path.join(db.data_dir(), "media", safe_name)
    import shutil
    try:
        shutil.copy2(src_path, dest)
    except OSError as ex:
        return {"ok": False, "error": f"Could not copy the document: {ex}"}

    try:
        text = extract_text(dest)
    except Exception as ex:
        try:
            os.remove(dest)
        except OSError:
            pass
        return {"ok": False, "error": f"Could not read that document: {ex}"}

    conn = db.connect()
    try:
        cur = conn.execute(
            "INSERT INTO documents (title, subject_id, grade, filename, original_name, extracted_text) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (title, subject_id, grade, safe_name, os.path.basename(src_path), text),
        )
        conn.commit()
        return {"ok": True, "id": cur.lastrowid}
    finally:
        conn.close()


def list_documents(grade: Optional[str] = None, subject_id: Optional[int] = None) -> List[Dict]:
    conn = db.connect()
    try:
        query = (
            "SELECT d.*, s.name AS subject_name FROM documents d "
            "LEFT JOIN subjects s ON s.id = d.subject_id WHERE 1=1"
        )
        params: List = []
        if grade:
            query += " AND d.grade = ?"
            params.append(grade)
        if subject_id:
            query += " AND d.subject_id = ?"
            params.append(subject_id)
        query += " ORDER BY d.created_at DESC"
        rows = conn.execute(query, params).fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


def delete_document(document_id: int) -> Dict:
    conn = db.connect()
    try:
        doc = conn.execute("SELECT * FROM documents WHERE id = ?", (document_id,)).fetchone()
        if not doc:
            return {"ok": False, "error": "Document not found."}
        conn.execute("DELETE FROM documents WHERE id = ?", (document_id,))
        conn.commit()
    finally:
        conn.close()
    try:
        os.remove(os.path.join(db.data_dir(), "media", doc["filename"]))
    except OSError:
        pass
    return {"ok": True}


def _fts_query(text: str) -> str:
    """FTS5 MATCH needs quoted terms to tolerate punctuation/apostrophes in a
    free-typed topic — build an OR-of-terms query rather than passing raw text."""
    terms = [t for t in text.replace('"', " ").split() if t]
    if not terms:
        return ""
    return " OR ".join(f'"{t}"' for t in terms)


def search_documents(grade: str, subject_id: Optional[int], query: str, limit: int = 3) -> List[str]:
    """Returns contextual snippets (not full documents) matching grade + subject
    + the teacher's topic, most relevant first. Empty list if nothing matches —
    callers should treat that as "no reference material available", not an error."""
    match = _fts_query(query)
    if not match:
        return []
    conn = db.connect()
    try:
        rows = conn.execute(
            """SELECT snippet(documents_fts, 1, '', '', '...', 40) AS excerpt
               FROM documents_fts
               JOIN documents ON documents.id = documents_fts.rowid
               WHERE documents.grade = ? AND documents.subject_id = ?
                 AND documents_fts MATCH ?
               ORDER BY rank LIMIT ?""",
            (grade, subject_id, match, limit),
        ).fetchall()
        return [r["excerpt"] for r in rows if r["excerpt"]]
    except Exception:
        # FTS5 MATCH syntax errors on odd input (e.g. lone operators) — degrade
        # to "no reference material" rather than breaking generation entirely.
        return []
    finally:
        conn.close()


def get_documents_text(document_ids: List[int], max_chars: int = 6000) -> List[str]:
    """Returns full (truncated) extracted_text for explicitly teacher-picked
    documents, in the order given — used instead of search_documents() when
    the teacher selects specific documents rather than relying on automatic
    keyword matching against the topic."""
    if not document_ids:
        return []
    conn = db.connect()
    try:
        placeholders = ",".join("?" * len(document_ids))
        rows = conn.execute(
            f"SELECT id, extracted_text FROM documents WHERE id IN ({placeholders})",
            document_ids,
        ).fetchall()
        by_id = {r["id"]: r["extracted_text"] for r in rows}
        return [by_id[i][:max_chars] for i in document_ids if by_id.get(i)]
    finally:
        conn.close()
